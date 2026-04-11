"""
Excel Export Module
Export detection results to Excel with formatting
"""
import pandas as pd
from typing import List, Dict
from pathlib import Path
import logging
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
import xlsxwriter

logger = logging.getLogger(__name__)

class ExcelExporter:
    """Export detection results to Excel"""
    
    def __init__(self):
        self.wb = None
        self.ws = None
    
    def export_to_excel(
        self,
        results: Dict,
        output_path: str,
        include_charts: bool = True,
        include_summary: bool = True
    ):
        """
        Export detection results to formatted Excel file
        
        Args:
            results: Detection results dictionary
            output_path: Path to save Excel file
            include_charts: Include visualization charts
            include_summary: Include summary sheet
        """
        logger.info(f"Exporting results to Excel: {output_path}")
        
        # Create main data sheet
        df_main = self._create_main_dataframe(results)
        
        # Export with xlsxwriter for better formatting
        with pd.ExcelWriter(output_path, engine='xlsxwriter') as writer:
            # Write main data
            df_main.to_excel(writer, sheet_name='Detections', index=False)
            
            # Get workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Detections']
            
            # Format main sheet
            self._format_main_sheet(workbook, worksheet, df_main)
            
            # Add summary sheet
            if include_summary:
                df_summary = self._create_summary_dataframe(results)
                df_summary.to_excel(writer, sheet_name='Summary', index=False)
                summary_ws = writer.sheets['Summary']
                self._format_summary_sheet(workbook, summary_ws, df_summary)
            
            # Add charts
            if include_charts:
                self._add_charts(workbook, worksheet, df_main, results)
        
        logger.info(f"✓ Excel file saved: {output_path}")
    
    def _create_main_dataframe(self, results: Dict) -> pd.DataFrame:
        """Create main dataframe with frame-by-frame counts"""
        
        target_classes = results['target_classes']
        frame_results = results['frame_results']
        
        # Initialize data structure
        data = []
        
        for frame_data in frame_results:
            row = {
                'Frame': frame_data['frame_number'],
                'Timestamp (s)': round(frame_data['timestamp'], 2),
                'Total Objects': sum(frame_data['counts'].values())
            }
            
            # Add count for each target class
            for cls in target_classes:
                row[cls.title()] = frame_data['counts'].get(cls, 0)
            
            data.append(row)
        
        df = pd.DataFrame(data)
        return df
    
    def _create_summary_dataframe(self, results: Dict) -> pd.DataFrame:
        """Create summary statistics dataframe"""
        
        stats = results['statistics']
        video_props = results['video_properties']
        
        summary_data = {
            'Metric': [
                'Video Duration (s)',
                'Total Frames',
                'Frames Processed',
                'FPS',
                'Resolution',
                'Total Detections',
                'Frames with Detections',
                'Detection Rate (%)'
            ],
            'Value': [
                round(video_props['duration'], 2),
                video_props['total_frames'],
                results['frames_processed'],
                round(video_props['fps'], 2),
                f"{video_props['width']}x{video_props['height']}",
                stats['total_detections'],
                stats['frames_with_detections'],
                round((stats['frames_with_detections'] / results['frames_processed']) * 100, 2)
            ]
        }
        
        df_summary = pd.DataFrame(summary_data)
        
        # Add class-specific statistics
        class_stats = []
        for cls in results['target_classes']:
            class_stats.append({
                'Object Class': cls.title(),
                'Total Count': stats['class_totals'][cls],
                'Average per Frame': round(stats['class_averages'][cls], 2),
                'Maximum in Single Frame': stats['class_max'][cls]
            })
        
        df_class_stats = pd.DataFrame(class_stats)
        
        return pd.concat([
            df_summary,
            pd.DataFrame([['', '']], columns=['Metric', 'Value']),
            pd.DataFrame([['Object Statistics', '']], columns=['Metric', 'Value']),
            df_class_stats.rename(columns={'Object Class': 'Metric', 'Total Count': 'Value'})
        ], ignore_index=True)
    
    def _format_main_sheet(self, workbook, worksheet, df):
        """Apply formatting to main detection sheet"""
        
        # Define formats
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#4472C4',
            'font_color': 'white',
            'align': 'center',
            'valign': 'vcenter',
            'border': 1
        })
        
        cell_format = workbook.add_format({
            'align': 'center',
            'valign': 'vcenter',
            'border': 1
        })
        
        number_format = workbook.add_format({
            'align': 'center',
            'valign': 'vcenter',
            'border': 1,
            'num_format': '0'
        })
        
        # Set column widths
        worksheet.set_column('A:A', 12)  # Frame
        worksheet.set_column('B:B', 15)  # Timestamp
        worksheet.set_column('C:C', 15)  # Total Objects
        worksheet.set_column('D:Z', 12)  # Object columns
        
        # Format header row
        for col_num, value in enumerate(df.columns.values):
            worksheet.write(0, col_num, value, header_format)
        
        # Set row height for header
        worksheet.set_row(0, 25)
    
    def _format_summary_sheet(self, workbook, worksheet, df):
        """Format summary sheet"""
        
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#70AD47',
            'font_color': 'white',
            'align': 'center',
            'border': 1
        })
        
        # Set column widths
        worksheet.set_column('A:A', 30)
        worksheet.set_column('B:Z', 20)
        
        # Format header
        for col_num, value in enumerate(df.columns.values):
            worksheet.write(0, col_num, value, header_format)
    
    def _add_charts(self, workbook, worksheet, df, results):
        """Add charts to Excel file"""
        
        # Create a new chart sheet
        chartsheet = workbook.add_chartsheet('Charts')
        
        # Line chart for object counts over time
        chart = workbook.add_chart({'type': 'line'})
        
        # Add data series for each object class
        for idx, cls in enumerate(results['target_classes'], start=4):
            col_letter = chr(65 + idx - 1)  # Convert to column letter
            chart.add_series({
                'name': f"=Detections!${col_letter}$1",
                'categories': f"=Detections!$A$2:$A${len(df) + 1}",
                'values': f"=Detections!${col_letter}$2:${col_letter}${len(df) + 1}",
            })
        
        chart.set_title({'name': 'Object Detection Over Time'})
        chart.set_x_axis({'name': 'Frame Number'})
        chart.set_y_axis({'name': 'Object Count'})
        chart.set_size({'width': 720, 'height': 480})
        
        chartsheet.set_chart(chart)

class CSVExporter:
    """Export results to CSV format"""
    
    @staticmethod
    def export_to_csv(results: Dict, output_path: str):
        """Export detection results to CSV"""
        
        target_classes = results['target_classes']
        frame_results = results['frame_results']
        
        data = []
        for frame_data in frame_results:
            row = {
                'Frame': frame_data['frame_number'],
                'Timestamp': frame_data['timestamp'],
                'Total': sum(frame_data['counts'].values())
            }
            
            for cls in target_classes:
                row[cls] = frame_data['counts'].get(cls, 0)
            
            data.append(row)
        
        df = pd.DataFrame(data)
        df.to_csv(output_path, index=False)
        
        logger.info(f"CSV exported: {output_path}")

class JSONExporter:
    """Export results to JSON format"""
    
    @staticmethod
    def export_to_json(results: Dict, output_path: str):
        """Export detection results to JSON"""
        import json
        
        with open(output_path, 'w') as f:
            json.dump(results, f, indent=2)
        
        logger.info(f"JSON exported: {output_path}")
